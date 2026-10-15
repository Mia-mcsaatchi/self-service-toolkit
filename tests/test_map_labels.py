"""Tests for scripts/map_labels.py against a fake Batches client (offline, free)."""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd

os.environ.setdefault("AUTH_DISABLED", "true")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import map_labels  # noqa: E402

CATS = {"Marvel": ["Marvel HYDRA organisation"], "Places and geography": ["Hydra / Ydra, Greek island"]}

ANSWERS = {  # label -> (broad, specific) the fake model returns
    "Marvel's Hydra organization": ("Marvel", "marvel hydra organisation"),     # case differs
    "HYDRA (Marvel organization)": ("Marvel", "Marvel HYDRA organisation"),
    "Greek island Hydra": ("Places and geography", "Hydra / Ydra, Greek island"),
    "roleplay character name": ("Marvel", "NEW: Roleplay character name"),
    "Hydra pub on Elm St": ("Places and geography", "Hydra pub"),                 # not in list -> NEW
    "unclear reference": ("Unclear", "Unspecified"),
    "hydrated skin": ("Not a Hydra reference", "Unspecified"),
}


class FakeBatches:
    def __init__(self, drop_first=False):
        self.batches, self.drop_first = {}, drop_first

    def create(self, requests):
        bid = f"b{len(self.batches)}"
        self.batches[bid] = requests
        return SimpleNamespace(id=bid)

    def retrieve(self, bid):
        return SimpleNamespace(id=bid, processing_status="ended")

    def results(self, bid):
        for req in self.batches[bid]:
            lines = req["params"]["messages"][0]["content"].split("\n")[1:]
            labels = [ln.split(". ", 1)[1] for ln in lines]
            items = [{"i": i, "broad": ANSWERS[lab][0], "specific": ANSWERS[lab][1]} for i, lab in enumerate(labels)]
            if self.drop_first:          # a partial answer must be resent, not half-applied
                items, self.drop_first = items[1:], False
            msg = SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps({"items": items}))],
                                  usage=SimpleNamespace(input_tokens=1000, output_tokens=200))
            yield SimpleNamespace(custom_id=req["custom_id"], result=SimpleNamespace(type="succeeded", message=msg))


def _client(b):
    return SimpleNamespace(messages=SimpleNamespace(batches=b))


def test_prompt_lists_taxonomy_and_schema_limits_broad():
    reqs = map_labels.build_requests(sorted(ANSWERS), CATS, "Hydra", "claude-sonnet-5")
    params, chunk = next(iter(reqs.values()))
    assert "Marvel: Marvel HYDRA organisation" in params["system"][0]["text"]
    enum = params["output_config"]["format"]["schema"]["properties"]["items"]["items"]["properties"]["broad"]["enum"]
    assert enum == ["Marvel", "Places and geography", "Unclear", "Not a Hydra reference"]
    assert len(chunk) == len(ANSWERS)


def test_run_maps_normalises_and_resends_partial_answers(tmp_path):
    batches = FakeBatches(drop_first=True)
    mapping, usage = map_labels.run(sorted(ANSWERS), CATS, "Hydra", _client(batches), "claude-sonnet-5",
                                    tmp_path / "s.json", poll_seconds=0)
    assert len(batches.batches) == 2           # first answer was incomplete, so it was resent
    assert mapping["Marvel's Hydra organization"] == ("Marvel", "Marvel HYDRA organisation")
    assert mapping["Hydra pub on Elm St"] == ("Places and geography", "NEW: Hydra pub")
    assert mapping["roleplay character name"] == ("Marvel", "NEW: Roleplay character name")
    assert usage["stopped"] is None


def test_breakdown_counts_exclusions_unresolved_and_new(tmp_path):
    mapping, _ = map_labels.run(sorted(ANSWERS), CATS, "Hydra", _client(FakeBatches()), "claude-sonnet-5",
                                tmp_path / "s.json", poll_seconds=0)
    df = pd.DataFrame({
        "context": ["Marvel's Hydra organization", "HYDRA (Marvel organization)", "Greek island Hydra",
                    "roleplay character name", "unclear reference", "hydrated skin", "Greek island Hydra"],
        "status": ["substantive", "incidental", "incidental", "incidental", "unclear", "incidental", "irrelevant"],
    })
    cat = map_labels.apply_mapping(df, mapping)
    assert list(cat["broad_category"]) == ["Marvel", "Marvel", "Places and geography", "Marvel",
                                           "Unresolved", "Excluded", "Excluded"]
    summary = map_labels.write_breakdown(cat, tmp_path / "b.xlsx")
    assert summary == {"total": 7, "excluded": 2, "population": 5, "unresolved": 1, "new_rows": 1, "new_tags": 1}
    wb = openpyxl.load_workbook(tmp_path / "b.xlsx")
    assert wb.sheetnames == ["HYDRA Associations", "Proposed new associations", "Label map"]
    cells = [r for r in wb["HYDRA Associations"].iter_rows(values_only=True)]
    assert ["Marvel", "Subtotal", 3] == list(cells[7][:3])
    assert any(r[1] == "Roleplay character name (new)" for r in cells)


def test_finalise_rules_merge_move_and_fold_small_tags():
    import finalise_breakdown
    cat = pd.DataFrame({
        "context": ["character name in roleplay"] * 3 + ["roleplay character name"] * 2
                   + ["Stargate Hydra bridge"] * 2 + ["LayerZero Stargate Hydra"] + ["tiny thing"],
        "broad_category": ["TV and stage"] * 3 + ["Unresolved"] * 2 + ["Crypto"] * 3 + ["Crypto"],
        "specific_association": ["Starlight Express character"] * 3 + ["Unspecified"] * 2
                                + ["Stargate Hydra bridge"] * 2 + ["LayerZero Stargate Hydra bridge", "Tiny"],
        "new_association": ["True"] * 3 + ["False"] * 2 + ["True"] * 4,
    })
    cfg = {"min_mentions": 2, "rules": [
        {"from": {"broad": "TV and stage", "specific": "Starlight Express character"},
         "to": {"broad": "Fiction", "specific": "Roleplay character", "new": True}},
        {"from": {"broad": "Unresolved", "label_regex": "role-?play"},
         "to": {"broad": "Fiction", "specific": "Roleplay character", "new": True}},
        {"from": {"specific": "Stargate Hydra bridge"},
         "to": {"broad": "Crypto", "specific": "LayerZero Stargate Hydra bridge", "new": True}},
    ]}
    out = finalise_breakdown.apply_rules(cat, cfg)
    assert (out["specific_association"] == "Roleplay character").sum() == 5
    assert (out["specific_association"] == "LayerZero Stargate Hydra bridge").sum() == 3
    assert out.loc[8, "specific_association"] == "Other (smaller associations)"   # 1 < min_mentions
    assert not out.loc[8, "new_association"]
